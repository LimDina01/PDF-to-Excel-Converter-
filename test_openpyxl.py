import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.workbook import Workbook

text = "Release Outward Check Deposit CHQ# 5,000.00 239481 To Chip Mong Commercial Bank Plc Drawer Name Original Amount 5,000.00 USD Deposit on Sep 29, 2026 10:25 AM REF# 122LH03262720012 REMARK: BY PUTH SINETH 0968202801 (\u1791\u17bc\u178f\u17cb \u17c1\u179b\u17b8\u1780\u1791\u17b8\u17e2)"
# Add a known illegal character to text just in case the original doesn't have one natively in string format
text_with_illegal = text + chr(1)

clean_text = ILLEGAL_CHARACTERS_RE.sub('', text_with_illegal)

wb = Workbook()
ws = wb.active

print('Trying original...')
try:
    ws['A1'] = text_with_illegal
    print('Original WORKED (this is unexpected if it has chr(1))')
except Exception as e:
    print('Original failed due to illegal chars.')

print('Trying cleaned...')
try:
    ws['A1'] = clean_text
    print('Cleaned worked!')
except Exception as e:
    print('Cleaned failed too!')
